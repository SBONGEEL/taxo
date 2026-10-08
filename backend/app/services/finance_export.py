"""تصديرُ الملخّصات المالية إلى Excel (SPEC §٦٥-د/٥) — **ملفُّ `.xlsx` حقيقيٌّ من الدوالّ نفسِها التي ترسم الشاشة**.

**ولا استعلامَ هنا**: كلُّ صفٍّ يُكتب جاء من `finance_summary.summary` أو `users_page` أو `transactions_page` — **البابُ
الذي يقرؤه العرض**. فلا يفترق رقمٌ في الملفّ عن رقمٍ في الشاشة: هما النداءُ نفسُه بصفحةٍ أطول.

**والمالُ رقمٌ لا نصّ** — `Decimal` يُكتب كما هو بصيغة `0.000` (لا float)، فيجمعه المحاسبُ في Excel ولا يقرأ «3.2499999».
**والوقتُ بيوم السوق** لا بـUTC: Excel لا يحمل منطقةً زمنية، ووقتٌ بلا منطقةٍ يُقرأ محليّاً — فيُكتب محليّاً.

**ووضعُ «الكتابة وحدها»** (`write_only`): الصفوفُ تُكتب ولا تُحفظ في الذاكرة كائناتٍ — والحدُّ `EXPORT_MAX_ROWS` يمنع ملفّاً
بلا قاع، **وما اقتُطع يُقال في آخر الورقة** لا يُسكت عنه.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from app.services import finance_summary as fs

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: **سقفُ الصفوف في ملفٍّ واحد** — عشرون ألفاً تكفي شهرَ سوقٍ بحجمه اليوم أضعافاً، **وما فوقها يُقتطع ويُقال**
EXPORT_MAX_ROWS = 20_000

_MONEY_FORMAT = "0.000"
_TIME_FORMAT = "yyyy-mm-dd hh:mm"

_PERIOD_LABEL = {"today": "اليوم", "week": "الأسبوع", "month": "الشهر", "custom": "مخصَّصة"}
_STATUS_FILTER_LABEL = {
    fs.FinanceStatus.CONFIRMED: "مؤكَّد",
    fs.FinanceStatus.UNCONFIRMED: "غيرُ مؤكَّد",
    fs.FinanceStatus.DISPUTED: "متنازَعٌ عليه",
}


class _Sheet:
    """ورقةٌ تُكتب من اليمين — **والمالُ والوقتُ بصيغتيهما** في موضعٍ واحدٍ لا في كلِّ سطر."""

    def __init__(self, book: Workbook, title: str, scope: fs.Scope) -> None:
        self.sheet = book.create_sheet(title)
        self.sheet.sheet_view.rightToLeft = True
        self.scope = scope
        self.bold = Font(bold=True)

    def header(self, *titles: str) -> None:
        cells = []
        for title in titles:
            cell = WriteOnlyCell(self.sheet, value=title)
            cell.font = self.bold
            cells.append(cell)
        self.sheet.append(cells)

    def row(self, *values: object) -> None:
        cells = []
        for value in values:
            if isinstance(value, Decimal):
                cell = WriteOnlyCell(self.sheet, value=value)
                cell.number_format = _MONEY_FORMAT
            elif isinstance(value, datetime):
                # **Excel لا يحمل منطقة** — فيُكتب الوقتُ بيوم السوق بلا منطقة، وهو ما يقرؤه صاحبُ السوق
                cell = WriteOnlyCell(self.sheet, value=value.astimezone(self.scope.zone).replace(tzinfo=None))
                cell.number_format = _TIME_FORMAT
            else:
                cell = WriteOnlyCell(self.sheet, value=value)
            cells.append(cell)
        self.sheet.append(cells)


def _about(book: Workbook, scope: fs.Scope, title: str) -> None:
    """ورقةُ «عن الملفّ» — **ما صُدِّر وبأيِّ مرشِّحات**، فلا يُقرأ ملفٌّ بعد شهرٍ بلا سياقه."""
    sheet = _Sheet(book, "عن الملفّ", scope)
    sheet.header("البند", "القيمة")
    sheet.row("العرض", title)
    sheet.row("السوق", scope.country.value)
    sheet.row("العملة", scope.currency.value)
    sheet.row("الفترة", _PERIOD_LABEL[scope.period])
    sheet.row("من (بيوم السوق)", scope.from_at)
    sheet.row("إلى — غيرُ داخل (بيوم السوق)", scope.to_at)
    sheet.row("المنطقة الزمنية", str(scope.zone))
    sheet.row("نوعُ المستخدم", fs.ROLE_LABEL.get(scope.user_type.value) if scope.user_type else "الكلّ")
    sheet.row("طريقةُ الدفع", fs.method_label(scope.method.value) if scope.method else "الكلّ")
    sheet.row("الحال", _STATUS_FILTER_LABEL[scope.status] if scope.status else "الكلّ")
    sheet.row("صُدِّر في", scope.now)
    sheet.row("حساباتُ التجربة", "مستثناةٌ دائماً")


def _finish(book: Workbook) -> bytes:
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def summary_workbook(result: fs.Summary) -> bytes:
    """**الملخّصُ كما في الشاشة**: كلُّ مجموعٍ بمجموعته ورقمه وعدده وتعريفه، وتفصيلُه تحته، ثمّ المطابقة."""
    scope = result.scope
    book = Workbook(write_only=True)
    _about(book, scope, "الملخّص")
    sheet = _Sheet(book, "الملخّص", scope)
    sheet.header("المجموعة", "البند", "المبلغ", "العملة", "العدد", "المستخدمون", "رصيدٌ الآن", "ينطبق", "التعريف")
    titles = {group.key: group.title for group in fs.GROUPS}
    for value in result.values:
        metric = value.metric
        sheet.row(
            titles[metric.group],
            metric.label,
            value.amount if value.applicable else None,
            scope.currency.value,
            value.count if value.applicable else None,
            value.users if value.applicable else None,
            "نعم" if metric.stock else "لا",
            "نعم" if value.applicable else (value.reason or "لا"),
            metric.definition,
        )
        for line in value.breakdown:
            sheet.row(titles[metric.group], f"{metric.label} · {line.label}", line.amount, scope.currency.value, line.count)

    check = _Sheet(book, "المطابقة", scope)
    reconciliation = result.reconciliation
    check.header("المطابقة", "الحال", "فُحص في", "عددُ الفحوص")
    check.row("المجاميعُ والدفتر", "تتطابق" if reconciliation.reconciled else "لا تتطابق", reconciliation.checked_at, reconciliation.checks)
    check.header("البند", "رقمُ الصفحة", "رقمُ الدفتر", "الفرق")
    for item in reconciliation.differences:
        check.row(item.label, item.page_amount, item.ledger_amount, item.difference_amount)
    return _finish(book)


def users_workbook(scope: fs.Scope, page: fs.Page) -> bytes:
    """**«المستخدمون» وراء مجموع** — الصفوفُ نفسُها التي ترسمها الصفحة، مرتَّبةً من الأكبر."""
    book = Workbook(write_only=True)
    metric = page.value.metric
    _about(book, scope, f"المستخدمون — {metric.label}")
    sheet = _Sheet(book, "المستخدمون", scope)
    sheet.header("الاسم", "الدور", "الهاتف", "عددُ المعاملات", "المجموع", "العملة", "معرّفُ الحساب")
    _not_applicable(sheet, page)
    for row in page.rows:
        sheet.row(row.name, row.role_label, row.phone_masked, row.count, row.amount, scope.currency.value, str(row.user_id) if row.user_id else None)
    _truncation(sheet, page)
    return _finish(book)


def transactions_workbook(scope: fs.Scope, page: fs.Page) -> bytes:
    """**«المعاملات» وراء مجموع** — الأكبرُ قيمةً أوّلاً كما في الصفحة."""
    book = Workbook(write_only=True)
    metric = page.value.metric
    _about(book, scope, f"المعاملات — {metric.label}")
    sheet = _Sheet(book, "المعاملات", scope)
    sheet.header("التاريخ", "المستخدم", "الدور", "النوع", "الطريقة", "الحال", "المبلغ", "العملة", "الرحلة", "المرجع")
    _not_applicable(sheet, page)
    for row in page.rows:
        sheet.row(
            row.occurred_at,
            row.name,
            row.role_label,
            row.kind_label,
            row.method_label,
            row.status_label,
            row.amount,
            scope.currency.value,
            str(row.ride_id) if row.ride_id else None,
            str(row.ref_id),
        )
    _truncation(sheet, page)
    return _finish(book)


def _not_applicable(sheet: _Sheet, page: fs.Page) -> None:
    """**مجموعٌ لا ينطبق عليه مرشِّحٌ يُقال بعلّته** كما تقوله الشاشة («لا ينطبق: …») — **لا ملفٌّ بعنوانه وحدَه** يُقرأ «لا
    معاملات في الفترة» وهو «لم يُسأل أصلاً»."""
    if not page.value.applicable:
        sheet.row(f"لا ينطبق: {page.value.reason}")


def _truncation(sheet: _Sheet, page: fs.Page) -> None:
    """**ما لم يُكتب يُقال** — صفٌّ أخيرٌ بالعدد الكامل حين يتجاوز السقف، لا ملفٌّ يُقرأ كاملاً وهو مقطوع."""
    if page.total > len(page.rows):
        sheet.row(f"اقتُطع الملفّ عند {len(page.rows)} من {page.total} — ضيّق الفترةَ أو المرشِّحات")
