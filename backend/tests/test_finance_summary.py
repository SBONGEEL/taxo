"""الملخّصاتُ المالية (SPEC §٦٥-د) — **الصلاحيةُ والتدقيقُ والسوقُ والفترةُ والتصديرُ والمطابقة**.

ومجاميعُ كلِّ بطاقةٍ بمساراتها الحقيقيّة في `test_finance_summary_totals.py`. **وهنا ما يحكم الصفحةَ كلَّها**:

- **الصلاحيةُ مطلوبة**: المشرفُ الكاملُ على افتراضه يُردّ ٤٠٣ — **ولا سطرَ تدقيقٍ لمن رُدّ**.
- **كلُّ عرضٍ وتصديرٍ سطرٌ في التدقيق** بالسوق والفترة والمرشِّحات وأيِّ عرض وأتصديرٌ هو.
- **حساباتُ التجربة خارجَ كلِّ رقم** — **والضابط**: بلا الوسم تعود الأرقامُ اثنين.
- **الأردنُ وليبيا لا يختلطان** — كلٌّ بعملته، ولا «كلّ الأسواق».
- **الفترةُ بيوم السوق** — بدالّةٍ محضةٍ بساعةٍ مثبَّتة، **وبالباب نفسِه**.
- **التصديرُ يساوي العرض** — يُفتح الملفُّ ويُقرأ رقماً رقماً.
- **المطابقةُ تصيح**: قيمةٌ مزوّرةٌ تُطفئ «تتطابق» وتسمّي الفرق.
"""

from __future__ import annotations

import uuid
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select, update

from app.models.audit import AdminAuditLog
from app.models.enums import CountryCode, FeatureKey, WalletTransactionType
from app.models.feature_flag import FeatureFlag
from app.models.notification import NotificationSetting
from app.models.user import User
from app.models.wallet_setting import WalletSetting
from app.services import finance_summary as fs
from tests.conftest import _staff_headers
from tests.helpers import (
    DRIVER,
    FAR_PICKUP,
    NEAR_PICKUP,
    OTHER_RIDER,
    RIDER,
    SECOND_DRIVER,
    approved_driver,
    bring_online,
    completed_ride,
    pay_ride,
    rider_session,
    topup_wallet,
)

# ═════════════════════════════════════════════════════════════ أدواتٌ يستعملها ملفُّ المجاميع أيضاً


async def finance_reader(client: AsyncClient, admin_headers: dict, phone: str = "+962790000071") -> dict:
    """**مشرفٌ مُنح «الملخّصات المالية» بالاسم** — كما تمنحها مصفوفةُ اللوحة (`Users.tsx::toggle`): مجموعتُه الفعّالةُ ومعها الجديدة."""
    headers = await _staff_headers("admin", phone, "قارئ الملخّصات")
    me = (await client.get("/auth/me", headers=headers)).json()
    rows = (await client.get("/admin/permissions", headers=admin_headers)).json()
    row = next(item for item in rows if item["user_id"] == me["id"])
    assert "finance.summary" not in row["permissions"], "الصلاحيةُ وصلت بالافتراض — وهي لا تُعطى إلا بالاسم"
    granted = await client.put(
        f"/admin/permissions/{me['id']}",
        json={"permissions": sorted(set(row["permissions"]) | {"finance.summary"})},
        headers=admin_headers,
    )
    assert granted.status_code == 200, granted.text
    return {"headers": headers, "id": me["id"]}


async def summary(client: AsyncClient, reader: dict, **params) -> dict:
    response = await client.get(
        "/admin/finance/summary",
        params={"country_code": "JO", "period": "month", **params},
        headers=reader["headers"],
    )
    assert response.status_code == 200, response.text
    return response.json()


def metric(body: dict, key: str) -> dict:
    for group in body["groups"]:
        for item in group["metrics"]:
            if item["key"] == key:
                return item
    raise AssertionError(f"لا مجموعَ باسم {key}")


def amount(body: dict, key: str) -> Decimal:
    item = metric(body, key)
    assert item["applicable"], f"{key} لا ينطبق: {item['reason']}"
    return Decimal(item["amount"])


async def drill(client: AsyncClient, reader: dict, key: str, view: str, **params) -> dict:
    response = await client.get(
        f"/admin/finance/summary/{key}/{view}",
        params={"country_code": "JO", "period": "month", "limit": 200, **params},
        headers=reader["headers"],
    )
    assert response.status_code == 200, response.text
    return response.json()


def assert_reconciled(body: dict) -> None:
    reconciliation = body["reconciliation"]
    assert reconciliation["reconciled"] is True, reconciliation["differences"]
    assert reconciliation["differences"] == []
    assert reconciliation["checks"] >= len(WalletTransactionType)


async def _audit_rows(session_factory) -> list[AdminAuditLog]:
    async with session_factory() as session:
        return list(
            (
                await session.scalars(
                    select(AdminAuditLog)
                    .where(AdminAuditLog.entity_type == "finance_summary")
                    .order_by(AdminAuditLog.created_at)
                )
            ).all()
        )


async def _mark_test(session_factory, user_id: str | uuid.UUID, value: bool = True) -> None:
    """الوسمُ كما يكتبه سكربتُ الإنشاء — **ولا بابَ في التطبيقات أو اللوحة يكتبه** (`test_accounts.py`)."""
    async with session_factory() as session:
        await session.execute(update(User).where(User.id == uuid.UUID(str(user_id))).values(is_test=value))
        await session.commit()


# ═════════════════════════════════════════════════════════════ ١) الصلاحية والتدقيق


async def test_a_default_admin_and_support_are_refused_and_nothing_is_audited(
    client: AsyncClient, admin_headers: dict, support_headers: dict, session_factory
) -> None:
    """**المشرفُ الكاملُ على افتراضه لا يملكها** — «لا تُعطى لكلِّ مشرفٍ افتراضاً». **والردُّ قبل أن يُقرأ رقم**: لا سطرَ تدقيق."""
    for headers in (admin_headers, support_headers):
        for path in (
            "/admin/finance/summary",
            "/admin/finance/summary/topups/users",
            "/admin/finance/summary/topups/transactions",
            "/admin/finance/summary/export",
        ):
            refused = await client.get(path, params={"country_code": "JO"}, headers=headers)
            assert refused.status_code == 403, (path, refused.text)
            assert refused.json()["code"] == "permission_denied"
    assert await _audit_rows(session_factory) == []

    # **ومن مُنحها يدخل** — والمنحُ نفسُه سطرٌ في تدقيق الصلاحيات لا في هذا
    reader = await finance_reader(client, admin_headers)
    assert (await client.get("/admin/finance/summary", params={"country_code": "JO"}, headers=reader["headers"])).status_code == 200


async def test_every_view_and_every_export_is_audited_with_its_filters(
    client: AsyncClient, admin_headers: dict, session_factory, jordan_wallet: None
) -> None:
    reader = await finance_reader(client, admin_headers)
    params = {"country_code": "JO", "period": "week", "user_type": "rider", "method": "cash", "status": "confirmed"}
    calls = [
        ("/admin/finance/summary", {}),
        ("/admin/finance/summary/topups/users", {"offset": 0}),
        ("/admin/finance/summary/topups/transactions", {"offset": 0}),
        ("/admin/finance/summary/export", {"view": "summary"}),
        ("/admin/finance/summary/export", {"view": "transactions", "metric": "topups"}),
    ]
    for path, extra in calls:
        response = await client.get(path, params=params | extra, headers=reader["headers"])
        assert response.status_code == 200, (path, response.text)

    rows = await _audit_rows(session_factory)
    assert len(rows) == len(calls)
    assert {str(row.actor_id) for row in rows} == {reader["id"]}
    assert all(row.action.value == "read" for row in rows)
    # **بلا ترتيب**: `created_at` لحظةُ بدء المعاملة في Postgres، **ونداءان متتاليان قد يتساويان أو يتبادلان** — والمقيسُ أن كلَّ
    # نداءٍ سطرُه بعينه، لا ترتيبُ الساعة
    seen = sorted(((row.details["view"], row.details["metric"], row.details["export"]) for row in rows), key=str)
    assert seen == sorted(
        [
            ("summary", None, False),
            ("users", "topups", False),
            ("transactions", "topups", False),
            ("summary", None, True),
            ("transactions", "topups", True),
        ],
        key=str,
    )
    for row in rows:
        assert row.details["country_code"] == "JO"
        assert row.details["period"] == "week"
        assert (row.details["user_type"], row.details["method"], row.details["status"]) == ("rider", "cash", "confirmed")
        # **والفترةُ بحدّيها** — من يقرأ السطرَ بعد شهرٍ يعرف ما رآه المشرف لا اسمَ الفترة وحدَه
        assert row.details["from_at"] < row.details["to_at"]


async def test_an_unknown_total_is_not_found_and_a_custom_period_needs_its_dates(
    client: AsyncClient, admin_headers: dict
) -> None:
    reader = await finance_reader(client, admin_headers)
    missing = await client.get(
        "/admin/finance/summary/not_a_total/users", params={"country_code": "JO"}, headers=reader["headers"]
    )
    assert missing.status_code == 404
    no_dates = await client.get(
        "/admin/finance/summary", params={"country_code": "JO", "period": "custom"}, headers=reader["headers"]
    )
    assert no_dates.status_code == 422 and no_dates.json()["code"] == "invalid_input"
    # **ولا «كلّ الأسواق»**: السوقُ إلزاميّ — دينارٌ أردنيٌّ لا يُجمع إلى ليبيّ
    no_market = await client.get("/admin/finance/summary", headers=reader["headers"])
    assert no_market.status_code == 422


# ═════════════════════════════════════════════════════════════ ٢) حساباتُ التجربة


async def test_test_accounts_are_left_out_of_every_total(
    client: AsyncClient, admin_headers: dict, session_factory, jordan_settings: None, jordan_wallet: None
) -> None:
    """رحلةٌ حقيقيّةٌ ورحلةُ تجربة، وتصحيحان، واشتراكان — **والأرقامُ كلُّها بواحد**؛ **والضابط**: بلا الوسم تعود اثنين."""
    reader = await finance_reader(client, admin_headers)
    real = await approved_driver(client, session_factory, DRIVER, plate_number="AMM-1001")
    tester = await approved_driver(client, session_factory, SECOND_DRIVER, plate_number="AMM-1002")
    await _mark_test(session_factory, tester["user_id"])
    await bring_online(client, real, NEAR_PICKUP)
    await bring_online(client, tester, FAR_PICKUP)

    real_rider = await rider_session(client)
    test_rider = await rider_session(client, OTHER_RIDER)
    await _mark_test(session_factory, test_rider["user"]["id"])

    for rider, driver in ((real_rider, real), (test_rider, tester)):
        ride = await completed_ride(client, rider["headers"], driver)
        paid = await pay_ride(client, rider["headers"], ride["id"], "cash", key=f"pay-{ride['id']}")
        assert paid.status_code == 201, paid.text
        confirmed = await client.post(f"/payments/{paid.json()['payments'][0]['id']}/confirm", headers=driver["headers"])
        assert confirmed.status_code == 200, confirmed.text

    # **تصحيحُ التجربة** — البابُ الوحيدُ الذي يعطي حسابَ التجربة رصيداً — **وتصحيحٌ حقيقيّ**
    for user_id, wallet, value in ((tester["user_id"], "driver", "40.000"), (real_rider["user"]["id"], "rider", "5.000")):
        done = await client.post(
            f"/admin/wallets/{user_id}/adjustments",
            json={"amount": value, "reason": "رصيدُ جولة", "wallet": wallet},
            headers=admin_headers,
        )
        assert done.status_code == 200, done.text

    body = await summary(client, reader)
    assert amount(body, "ride_payments_confirmed") == Decimal("8.000")
    assert metric(body, "ride_payments_confirmed")["count"] == 1
    assert amount(body, "adjustments") == Decimal("5.000")
    assert amount(body, "rider_balances") == Decimal("5.000")
    assert amount(body, "captain_balances") == Decimal("0.000")
    assert amount(body, "subscriptions_sold") == Decimal("30.000")
    assert metric(body, "subscriptions_sold")["count"] == 1
    # **والمطابقةُ تستثنيهم كذلك** — الدفترُ الخامُ لا يرى التجربةَ أيضاً، فلا فرقَ يُخترع
    assert_reconciled(body)
    users = await drill(client, reader, "ride_payments_confirmed", "users")
    assert [row["user_id"] for row in users["rows"]] == [real["user_id"]]

    # **والضابط** — بلا الوسم تعود الأرقامُ اثنين، فالواحدُ فوق أثرُ الاستثناء لا قلّةُ الصفوف
    await _mark_test(session_factory, tester["user_id"], False)
    await _mark_test(session_factory, test_rider["user"]["id"], False)
    again = await summary(client, reader)
    assert amount(again, "ride_payments_confirmed") == Decimal("16.000")
    assert amount(again, "adjustments") == Decimal("45.000")
    assert amount(again, "captain_balances") == Decimal("40.000")
    assert amount(again, "subscriptions_sold") == Decimal("60.000")
    assert_reconciled(again)


# ═════════════════════════════════════════════════════════════ ٣) الأسواق


async def test_jordan_and_libya_never_mix_and_each_carries_its_currency(
    client: AsyncClient, admin_headers: dict, session_factory, jordan_wallet: None
) -> None:
    async with session_factory() as session:
        session.add(FeatureFlag(country_code=CountryCode.LY, feature_key=FeatureKey.WALLET_ENABLED.value, enabled=True))
        session.add(WalletSetting(country_code=CountryCode.LY))
        await session.commit()
    reader = await finance_reader(client, admin_headers)
    jordan = await rider_session(client)
    libya = await rider_session(client, RIDER | {"phone": "0911234567", "name": "راكبٌ في طرابلس", "country_code": "LY"})
    await topup_wallet(client, admin_headers, jordan["user"]["id"], "10.000")
    await topup_wallet(client, admin_headers, libya["user"]["id"], "7.000")

    jo = await summary(client, reader, country_code="JO")
    ly = await summary(client, reader, country_code="LY")
    assert (jo["window"]["currency"], ly["window"]["currency"]) == ("JOD", "LYD")
    assert (amount(jo, "rider_balances"), amount(ly, "rider_balances")) == (Decimal("10.000"), Decimal("7.000"))
    assert (amount(jo, "topups"), amount(ly, "topups")) == (Decimal("10.000"), Decimal("7.000"))
    assert (amount(jo, "ledger.topup"), amount(ly, "ledger.topup")) == (Decimal("10.000"), Decimal("7.000"))
    assert_reconciled(jo)
    assert_reconciled(ly)

    # **وما وراء البطاقة من سوقها وحدَه**
    jo_users = await drill(client, reader, "topups", "users", country_code="JO")
    ly_users = await drill(client, reader, "topups", "users", country_code="LY")
    assert [row["user_id"] for row in jo_users["rows"]] == [jordan["user"]["id"]]
    assert [row["user_id"] for row in ly_users["rows"]] == [libya["user"]["id"]]
    assert ly_users["window"]["currency"] == "LYD"


# ═════════════════════════════════════════════════════════════ ٤) يومُ السوق


def test_the_window_is_the_markets_day_not_the_servers() -> None:
    """**الساعةُ مثبَّتة**: ٢٢:٣٠ UTC يومَ ٨ — **وهي ٠١:٣٠ من يوم ٩ في عمّان و٠٠:٣٠ منه في طرابلس**. فـ«اليوم» يومُ ٩ في السوقين،
    **وخادمٌ يقرأ يومَ UTC كان سيبدأه منتصفَ ليل ٨ بتوقيت غرينتش** ويضمّ ٢١ ساعةً من أمس السوق."""
    now = datetime(2026, 10, 8, 22, 30, tzinfo=UTC)
    amman = fs.window(ZoneInfo("Asia/Amman"), "today", now, None, None)
    assert amman == (datetime(2026, 10, 8, 21, 0, tzinfo=UTC), datetime(2026, 10, 9, 21, 0, tzinfo=UTC))
    tripoli = fs.window(ZoneInfo("Africa/Tripoli"), "today", now, None, None)
    assert tripoli == (datetime(2026, 10, 8, 22, 0, tzinfo=UTC), datetime(2026, 10, 9, 22, 0, tzinfo=UTC))

    week = fs.window(ZoneInfo("Asia/Amman"), "week", now, None, None)
    assert week[0] == datetime(2026, 10, 2, 21, 0, tzinfo=UTC)
    custom = fs.window(ZoneInfo("Asia/Amman"), "custom", now, date(2026, 10, 1), date(2026, 10, 9))
    assert custom == (datetime(2026, 9, 30, 21, 0, tzinfo=UTC), datetime(2026, 10, 9, 21, 0, tzinfo=UTC))


async def test_the_period_boundary_is_read_at_the_markets_midnight(
    client: AsyncClient, admin_headers: dict, session_factory, jordan_wallet: None
) -> None:
    """**بالباب نفسِه**: شحنةٌ وقعت الآن تدخل يومَها بتوقيت عمّان ولا تدخل أمسَه ولا غدَه — **وحدُّ النافذة منتصفُ ليل عمّان**."""
    async with session_factory() as session:
        session.add(
            NotificationSetting(
                country_code=CountryCode.LY,
                quiet_hours_start=datetime(2026, 1, 1, 22).time(),
                quiet_hours_end=datetime(2026, 1, 1, 8).time(),
                timezone="Africa/Tripoli",
            )
        )
        await session.commit()
    reader = await finance_reader(client, admin_headers)
    rider = await rider_session(client)
    await topup_wallet(client, admin_headers, rider["user"]["id"], "10.000")

    amman = ZoneInfo("Asia/Amman")
    local_day = datetime.now(amman).date()
    for day, expected in (
        (local_day, Decimal("10.000")),
        (local_day - timedelta(days=1), Decimal("0.000")),
        (local_day + timedelta(days=1), Decimal("0.000")),
    ):
        body = await summary(client, reader, period="custom", from_date=day.isoformat(), to_date=day.isoformat())
        assert amount(body, "topups") == expected, day
        start = datetime.fromisoformat(body["window"]["from_at"])
        assert start == datetime.combine(day, datetime.min.time(), tzinfo=amman)
        assert body["window"]["timezone"] == "Asia/Amman"

    today = await summary(client, reader, period="today")
    assert datetime.fromisoformat(today["window"]["from_at"]) == datetime.combine(
        local_day, datetime.min.time(), tzinfo=amman
    )
    # **وليبيا بمِنطقتها** — منتصفُ ليل طرابلس لا عمّان
    tripoli = ZoneInfo("Africa/Tripoli")
    libya = await summary(client, reader, country_code="LY", period="today")
    assert libya["window"]["timezone"] == "Africa/Tripoli"
    assert datetime.fromisoformat(libya["window"]["from_at"]) == datetime.combine(
        datetime.now(tripoli).date(), datetime.min.time(), tzinfo=tripoli
    )


# ═════════════════════════════════════════════════════════════ ٥) التصدير


def _sheet_rows(content: bytes, title: str) -> list[tuple]:
    book = load_workbook(BytesIO(content))
    assert book[title].sheet_view.rightToLeft is True
    return [tuple(row) for row in book[title].iter_rows(values_only=True)]


def _money(value: object) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.001"))


async def test_the_export_is_a_real_xlsx_and_equals_the_view(
    client: AsyncClient, admin_headers: dict, session_factory, jordan_wallet: None
) -> None:
    reader = await finance_reader(client, admin_headers)
    first = await rider_session(client)
    second = await rider_session(client, OTHER_RIDER)
    for rider, value in ((first, "12.500"), (second, "7.250"), (first, "3.125")):
        await topup_wallet(client, admin_headers, rider["user"]["id"], value)

    params = {"country_code": "JO", "period": "month"}
    view = await summary(client, reader)
    exported = await client.get("/admin/finance/summary/export", params=params, headers=reader["headers"])
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert exported.headers["content-disposition"].startswith("attachment;")
    assert exported.content[:2] == b"PK", "ليس ملفَّ xlsx — الملفُّ حزمةُ zip"

    titles = {group["key"]: group["title"] for group in view["groups"]}
    rows = _sheet_rows(exported.content, "الملخّص")
    by_label = {(row[0], row[1]): row for row in rows[1:]}
    compared = 0
    for group in view["groups"]:
        for item in group["metrics"]:
            row = by_label[(titles[group["key"]], item["label"])]
            assert _money(row[2]) == Decimal(item["amount"]), item["key"]
            assert row[4] == item["count"] and row[5] == item["users"], item["key"]
            compared += 1
    assert compared == len(fs.METRICS)

    # **ما وراء البطاقة بالملفّ نفسِه** — الصفوفُ نفسُها بترتيبها
    for view_name, sheet in (("users", "المستخدمون"), ("transactions", "المعاملات")):
        page = await drill(client, reader, "topups", view_name)
        out = await client.get(
            "/admin/finance/summary/export",
            params=params | {"view": view_name, "metric": "topups"},
            headers=reader["headers"],
        )
        assert out.status_code == 200, out.text
        sheet_rows = _sheet_rows(out.content, sheet)[1:]
        assert len(sheet_rows) == len(page["rows"]) == page["total"]
        if view_name == "users":
            assert [(row[6], _money(row[4]), row[3]) for row in sheet_rows] == [
                (item["user_id"], Decimal(item["amount"]), item["count"]) for item in page["rows"]
            ]
        else:
            assert [(row[9], _money(row[6])) for row in sheet_rows] == [
                (item["ref_id"], Decimal(item["amount"])) for item in page["rows"]
            ]
    assert [Decimal(row["amount"]) for row in (await drill(client, reader, "topups", "users"))["rows"]] == [
        Decimal("15.625"),
        Decimal("7.250"),
    ]

    # **ومجموعٌ لا ينطبق عليه المرشِّحُ يقول علّتَه في الملفّ كما في الشاشة** — لا ورقةٌ بعنوانها وحدَه تُقرأ «لا معاملات»
    for view_name, sheet in (("users", "المستخدمون"), ("transactions", "المعاملات")):
        screen = await drill(client, reader, "rider_balances", view_name, user_type="driver")
        assert screen["metric"]["applicable"] is False and screen["rows"] == []
        out = await client.get(
            "/admin/finance/summary/export",
            params=params | {"view": view_name, "metric": "rider_balances", "user_type": "driver"},
            headers=reader["headers"],
        )
        assert out.status_code == 200, out.text
        body = _sheet_rows(out.content, sheet)[1:]
        assert [row[0] for row in body] == [f"لا ينطبق: {screen['metric']['reason']}"]
        assert all(cell is None for cell in body[0][1:])


# ═════════════════════════════════════════════════════════════ ٦) المطابقة


def test_the_comparison_flags_a_tampered_total() -> None:
    """**القيمةُ المزوّرةُ تُطفئ «تتطابق» وتسمّي الفرق** — والمقارنةُ دالّةٌ محضةٌ فلا تحتاج قاعدة."""
    now = datetime(2026, 10, 8, 10, 42, tzinfo=UTC)
    wallets = {"rider": Decimal("40.000"), "driver": Decimal("12.200")}
    by_type = {"topup": Decimal("40.000"), "commission": Decimal("-0.800")}
    page = {
        "rider_balances": Decimal("40.000"),
        "captain_balances": Decimal("12.200"),
        "captain_balances_cliq": Decimal("0"),
        "captain_balances_other": Decimal("12.200"),
        "captain_payable_now": Decimal("12.200"),
        "captain_not_due": Decimal("0"),
        "topups": Decimal("40.000"),
        "topups_cash": Decimal("40.000"),
        "commission_collected": Decimal("0.800"),
        "ledger.topup": Decimal("40.000"),
        "ledger.commission": Decimal("-0.800"),
    }
    raw_split = {"payable": Decimal("12.200"), "cliq": Decimal("0")}
    clean = fs.compare(fs.checks_for(page, wallets, by_type, raw_split), checked_at=now)
    assert clean.reconciled is True and clean.differences == []

    tampered = fs.compare(
        fs.checks_for(page | {"rider_balances": Decimal("43.250")}, wallets, by_type, raw_split), checked_at=now
    )
    assert tampered.reconciled is False
    assert [(d.key, d.page_amount, d.ledger_amount, d.difference_amount) for d in tampered.differences] == [
        ("rider_wallets", Decimal("43.250"), Decimal("40.000"), Decimal("3.250"))
    ]
    # **وقسمةٌ لا تجمع إلى أصلها تصيح أيضاً** — قناةٌ ضاع منها شحن
    split = fs.compare(fs.checks_for(page | {"topups_cash": Decimal("39.000")}, wallets, by_type, raw_split), checked_at=now)
    assert [d.key for d in split.differences] == ["topups_by_channel"]

    # **وقسمةٌ خاطئةٌ تجمع إلى أصلها تصيح كذلك** — كبتنٌ مجمَّدٌ عُدّ رصيدُه «يُصرف الآن»، وكبتنٌ بلا كليك عُدّ في «كليك»: **الجمعُ
    # يساوي الأصلَ في الحالين**، وفحصُ «القسمان يجمعان إلى الأصل» كان يُخضرّهما. **والقرينُ الخامُ يمسكهما**
    misplaced = page | {
        "captain_payable_now": Decimal("10.000"),
        "captain_not_due": Decimal("2.200"),
        "captain_balances_cliq": Decimal("4.000"),
        "captain_balances_other": Decimal("8.200"),
    }
    wrong = fs.compare(fs.checks_for(misplaced, wallets, by_type, raw_split), checked_at=now)
    assert [(d.key, d.page_amount, d.ledger_amount) for d in wrong.differences] == [
        ("driver_wallets_cliq", Decimal("4.000"), Decimal("0.000")),
        ("driver_wallets_other", Decimal("8.200"), Decimal("12.200")),
        ("driver_wallets_payable", Decimal("10.000"), Decimal("12.200")),
        ("driver_wallets_not_due", Decimal("2.200"), Decimal("0.000")),
    ]


async def test_the_page_carries_the_red_warning_when_the_ledger_disagrees(
    client: AsyncClient, admin_headers: dict, jordan_wallet: None, monkeypatch
) -> None:
    """**وبالباب**: دفترٌ «يقول» غيرَ ما تقوله الصفحة ⇒ `reconciled: false` بفرقه — **وهو ما ترسمه اللوحةُ شريطاً أحمر**."""
    reader = await finance_reader(client, admin_headers)
    rider = await rider_session(client)
    await topup_wallet(client, admin_headers, rider["user"]["id"], "10.000")
    assert_reconciled(await summary(client, reader))

    real = fs._raw_ledger

    async def drifted(session, scope):
        wallets, by_type = await real(session, scope)
        return wallets | {"rider": wallets["rider"] + Decimal("3.250")}, by_type

    monkeypatch.setattr(fs, "_raw_ledger", drifted)
    body = await summary(client, reader)
    reconciliation = body["reconciliation"]
    assert reconciliation["reconciled"] is False
    assert [
        (d["key"], d["page_amount"], d["ledger_amount"], d["difference_amount"]) for d in reconciliation["differences"]
    ] == [("rider_wallets", "10.000", "13.250", "-3.250")]


def test_every_ledger_type_has_a_name_and_a_reconciliation_line() -> None:
    """**نوعُ قيدٍ يُضاف إلى التعداد بلا اسمٍ هنا يُرسم بمفتاحه في الصفحة والملفّ** — فيسقط هذا قبل أن يُرى."""
    assert set(fs.LEDGER_LABEL) == {kind.value for kind in WalletTransactionType}
    keys = {check.key for check in fs.checks_for({}, {}, {}, {})}
    assert {f"ledger.{kind.value}" for kind in WalletTransactionType} <= keys
    assert {f"ledger.{kind.value}" for kind in WalletTransactionType} <= set(fs.METRICS)
